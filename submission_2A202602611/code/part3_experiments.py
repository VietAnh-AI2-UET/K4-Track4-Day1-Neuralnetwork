"""Các phép so sánh Part 3; chỉ dùng train và validation để chọn cấu hình."""
from __future__ import annotations

import json
import math
from pathlib import Path

import numpy as np
import pandas as pd
import torch

from data import make_val_split, fit_standardizer, apply_standardizer
from plots import plot_run
from results_table import load_results, save_result, to_row, write_xlsx
from train import run_experiment


TRIALS = [
    dict(exp_id="opt-adam-lr1e-3-s1", group="optimizer",
         description="Adam lr=0.001; so với SGD+momentum ở lr tốt nhất",
         changes={"optimizer": "adam", "lr": 0.001},
         prediction="Adam với lr=0.001 có thể giảm val loss nhanh hơn baseline trong ba epoch đầu, "
                    "vì nó điều chỉnh mức cập nhật riêng cho từng trọng số. Tôi kiểm tra tốc độ giảm loss, "
                    "macro-F1 tại epoch có val loss thấp nhất và độ lớn gradient.",
         notes="Đổi optimizer và lr so với baseline để chọn lr phù hợp cho Adam; "
               "momentum không áp dụng cho Adam; betas=(0.9,0.999); eps=1e-8; weight_decay=0"),
    dict(exp_id="opt-adam-lr3e-3-s1", group="optimizer",
         description="Adam lr=0.003; khảo sát lr thứ hai cho Adam",
         changes={"optimizer": "adam", "lr": 0.003},
         prediction="Tăng lr của Adam từ 0.001 lên 0.003 có thể giúp giảm loss nhanh hơn ở đầu quá trình, "
                    "nhưng cũng có thể làm loss và gradient dao động hơn. Tôi so hai lần chạy Adam và "
                    "chọn lr bằng macro-F1 tại epoch có val loss thấp nhất.",
         notes="Đổi optimizer và lr so với baseline; so với opt-adam-lr1e-3-s1 chỉ đổi lr; "
               "momentum không áp dụng cho Adam; betas=(0.9,0.999); eps=1e-8; weight_decay=0"),
    dict(exp_id="hp-batch128-s1", group="hparam",
         description="Chỉ giảm batch từ 512 xuống 128; giữ lr=0.1",
         changes={"batch": 128},
         prediction="Batch 128 tạo gần bốn lần số bước cập nhật so với batch 512 trong cùng 20 epoch; "
                    "model có thể đạt val tốt hơn nhưng mất nhiều thời gian và dao động hơn ở lr=0.1. "
                    "Tôi kiểm tra số bước, thời gian, đường loss và macro-F1.",
         notes="Chỉ đổi batch; lr không tăng/giảm theo batch; giữ lô cuối; "
               "cùng epoch nhưng khác số bước nên không thể quy mọi khác biệt cho kích thước batch"),
    dict(exp_id="hp-batch2048-s1", group="hparam",
         description="Chỉ tăng batch từ 512 lên 2048; giữ lr=0.1",
         changes={"batch": 2048},
         prediction="Batch 2048 chỉ có khoảng một phần tư số bước cập nhật của baseline, nên với cùng "
                    "lr=0.1 model có thể học chậm hơn tính theo epoch và đạt macro-F1 thấp hơn. "
                    "Thời gian có thể giảm; tôi chỉ kết luận sau khi đo thời gian và số bước.",
         notes="Chỉ đổi batch; lr không thay đổi; giữ lô cuối; chưa thử tăng lr theo batch "
               "hoặc giữ cùng tổng số bước cập nhật"),
]


def prepare_part3_data(processed_dir: str, device: str) -> dict:
    """Nạp riêng train; chia và chuẩn hoá giống Part 0, không nạp eval."""
    with np.load(Path(processed_dir) / "train.npz", allow_pickle=False) as source:
        X, y = source["X"], source["y"]
    X_tr, y_tr, X_val, y_val = make_val_split(X, y, val_fraction=0.2, seed=42)
    mean, std = fit_standardizer(X_tr)
    return {
        "X_tr": torch.tensor(apply_standardizer(X_tr, mean, std), device=device),
        "y_tr": torch.tensor(y_tr, device=device),
        "X_val": torch.tensor(apply_standardizer(X_val, mean, std), device=device),
        "y_val": torch.tensor(y_val, device=device),
    }


def run_trial(baseline_cfg, trial, data, output_dir):
    cfg = {**baseline_cfg, **trial["changes"], **{
        key: trial[key] for key in ("exp_id", "group", "description", "notes")
    }, "betas": (0.9, 0.999), "eps": 1e-8, "verbose": True}
    # Không để thay đổi ngoài kế hoạch lọt vào thí nghiệm.
    for key in ("loss", "optimizer", "lr", "weight_decay", "momentum", "batch",
                "epochs", "hidden", "dropout", "init", "clip_norm", "precision", "seed"):
        if key not in trial["changes"]:
            assert cfg[key] == baseline_cfg[key], f"Unexpected change: {key}"
    print(json.dumps(cfg, ensure_ascii=False, indent=2))
    result = run_experiment(cfg, data)
    output_dir = Path(output_dir)
    save_result(result, str(output_dir / "results"))
    plot_run(result, str(output_dir / "figures" / f"{cfg['exp_id']}.png"))
    print(json.dumps(result["summary"], ensure_ascii=False, indent=2))
    return result


def compare_text(result, reference, noise_f1, noise_acc, n_train, prediction_kind):
    """Nhận xét dùng số vừa đo; tách quan sát khỏi cách giải thích có thể phù hợp."""
    s, b = result["summary"], reference["summary"]
    h, bh = result["history"], reference["history"]
    if s["diverged"] or s["best_epoch"] <= 0:
        return (f"**Đối chiếu:** lần chạy phân kỳ hoặc chưa có epoch hợp lệ: "
                f"{s.get('divergence_reason', 'xem JSON')}. Giữ kết quả và cờ lỗi; không chọn làm cấu hình cuối.")
    df1, dacc = s["val_macro_f1"] - b["val_macro_f1"], s["val_acc"] - b["val_acc"]
    noise = ("vượt" if abs(df1) > noise_f1 else "chưa vượt")
    acc_noise = "vượt" if abs(dacc) > noise_acc else "chưa vượt"
    epoch3 = min(3, len(h["epoch"]), len(bh["epoch"])) - 1
    faster_early = h["val_loss"][epoch3] < bh["val_loss"][epoch3]
    parts = [
        f"**Đối chiếu sau khi chạy:** `{result['cfg']['exp_id']}` chọn epoch **{s['best_epoch']}** "
        f"với val loss **{s['best_val_loss']:.6f}**, accuracy **{s['val_acc']:.2%}**, "
        f"macro-F1 **{s['val_macro_f1']:.6f}**. So với `{reference['cfg']['exp_id']}`, "
        f"F1 đổi **{df1:+.6f}**, {noise} ngưỡng nhiễu **{noise_f1:.6f}**; "
        f"accuracy đổi **{dacc * 100:+.3f} điểm phần trăm**, {acc_noise} ngưỡng "
        f"**{noise_acc * 100:.3f} điểm phần trăm**.",
    ]
    if prediction_kind == "adam":
        parts.append(
            f"Dự đoán giảm loss nhanh trong ba epoch đầu {'khớp' if faster_early else 'chưa khớp'} "
            f"khi so với baseline: tại epoch {epoch3 + 1}, val loss là "
            f"**{h['val_loss'][epoch3]:.4f}** so với **{bh['val_loss'][epoch3]:.4f}**. "
            "Adam điều chỉnh mức cập nhật riêng theo lịch sử gradient của từng trọng số; "
            "cách giải thích này phù hợp với chênh lệch đường cong, nhưng lần chạy này chưa tách riêng "
            "ảnh hưởng của optimizer khỏi learning rate."
        )
    else:
        steps = s["total_update_steps"]
        base_steps = math.ceil(n_train / reference["cfg"]["batch"]) * len(bh["epoch"])
        total_time = s["total_time_s"]
        base_time = sum(bh["epoch_time_s"])
        expected_direction = df1 > 0 if prediction_kind == "small_batch" else df1 < 0
        parts.append(
            f"Dự đoán về hướng thay đổi F1 {'khớp' if expected_direction else 'khác kết quả'}. "
            f"Số bước cập nhật đo được là **{steps:,}**, baseline có **{base_steps:,}** bước "
            "(tính từ số mẫu, batch và số epoch hoàn tất). "
            f"Tổng thời gian epoch đo được **{total_time:.1f}s** so với baseline đã lưu **{base_time:.1f}s**. "
            "Batch nhỏ cập nhật nhiều lần hơn, batch lớn cập nhật ít lần hơn; vì vậy so cùng epoch "
            "chưa tách được ảnh hưởng của batch khỏi số bước. Thời gian baseline là số đo cũ, "
            "không có log số luồng CPU nên chỉ tham khảo, chưa đủ kết luận chắc chắn về tốc độ."
        )
    parts.append(
        f"Ở epoch cuối, train loss **{h['train_loss'][-1]:.4f}**, val loss **{h['val_loss'][-1]:.4f}**, "
        f"khoảng cách **{h['val_loss'][-1] - h['train_loss'][-1]:+.4f}**; cả hai được đo ở `eval()`. "
        f"Gradient trung bình theo epoch nằm trong **{min(h['grad_norm']):.4f}–{max(h['grad_norm']):.4f}** "
        "(trước clip; thí nghiệm này không clip). Một điểm loss hoặc gradient chưa đủ kết luận về độ ổn định."
    )
    parts.append("Ngưỡng nhiễu chỉ ước lượng từ ba seed baseline; cấu hình mới mới chạy seed 1. "
                 "Đây là quan sát trên validation, chưa chứng minh cấu hình thắng ổn định khi đổi seed.")
    return "\n\n".join(parts)


def comparison_table(results, n_train):
    return pd.DataFrame([{
        "exp_id": r["cfg"]["exp_id"], "optimizer": r["cfg"]["optimizer"],
        "lr": r["cfg"]["lr"], "batch": r["cfg"]["batch"],
        "best_epoch": r["summary"]["best_epoch"],
        "best_val_loss": r["summary"]["best_val_loss"],
        "val_acc": r["summary"]["val_acc"], "val_macro_f1": r["summary"]["val_macro_f1"],
        "update_steps": r["summary"].get("total_update_steps",
            math.ceil(n_train / r["cfg"]["batch"]) * len(r["history"]["epoch"])),
        "time_per_epoch_s": r["summary"]["time_per_epoch_s"],
        "total_time_s": sum(r["history"]["epoch_time_s"]),
        "diverged": r["summary"]["diverged"],
    } for r in results])


def select_by_val(results):
    valid = [r for r in results if not r["summary"]["diverged"]
             and r["summary"]["best_epoch"] > 0
             and np.isfinite(r["summary"]["val_macro_f1"])]
    if not valid:
        raise RuntimeError("No valid configuration to select")
    return max(valid, key=lambda r: (r["summary"]["val_macro_f1"],
                                    -r["summary"]["best_val_loss"], -r["cfg"]["lr"]))


def export_part3_table(output_dir, template_path, selected, summary_notes):
    import openpyxl

    output_dir = Path(output_dir)
    all_results = load_results(str(output_dir / "results"))
    rows = []
    for result in all_results:
        notes = ("Được chọn làm cấu hình cuối ở Part 3 bằng validation; chưa đánh giá eval"
                 if result["cfg"]["exp_id"] == selected["cfg"]["exp_id"] else "Chưa đánh giá eval")
        rows.append(to_row(result, notes=notes))
    out_path = output_dir / "experiments.xlsx"
    write_xlsx(rows, str(template_path), str(out_path))
    wb = openpyxl.load_workbook(out_path)
    ws = wb["Summary"]
    for index in range(2, 12):
        group = ws.cell(index, 1).value
        existing = [r["cfg"]["exp_id"] for r in all_results if r["cfg"]["group"] == group]
        note = summary_notes.get(group)
        if note is None:
            note = ("Log đã có ngoài hai chủ đề của phần hoàn thiện này: " + ", ".join(existing)
                    if existing else "Chưa khảo sát trong phạm vi Part 3 này")
        if group == "final" and existing:
            note += " Log kiểm tra seed đã có được giữ nguyên: " + ", ".join(existing) + "."
        ws.cell(index, 8, note)
    wb.save(out_path)
    return out_path
