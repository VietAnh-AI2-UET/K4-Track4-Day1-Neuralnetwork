"""Part 4: giữ lựa chọn trên val, chấm bằng script chính thức, xuất bằng chứng."""
from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path

import numpy as np
import torch

from train import DEFAULT_CFG, evaluate, final_eval, run_experiment
from results_table import load_results, save_result, to_row, write_xlsx


def recover_result(saved, data, checkpoint_dir):
    """Chạy lại đúng cấu hình khi JSON không có trọng số; kiểm tra trên val."""
    if "best_state" in saved:
        return saved
    cfg = {**DEFAULT_CFG, **saved["cfg"]}
    path = Path(checkpoint_dir) / f"{cfg['exp_id']}.pt"
    if path.exists():
        # TorchVersion do PyTorch tạo; vẫn nạp theo weights_only, không dùng pickle tuỳ ý.
        with torch.serialization.safe_globals([torch.torch_version.TorchVersion]):
            checkpoint = torch.load(path, map_location="cpu", weights_only=True)
        if checkpoint["cfg"] != cfg:
            raise ValueError(f"Checkpoint does not match selected run: {path}")
        result = {**checkpoint["result"], "best_state": checkpoint["best_state"]}
        from model import MLP
        model = MLP(hidden=tuple(cfg["hidden"]), dropout=cfg["dropout"], init=cfg["init"]).to(data["X_val"].device)
        model.load_state_dict(result["best_state"])
        metrics = evaluate(model, data["X_val"], data["y_val"], cfg["loss"])
        if not np.isclose(metrics["loss"], result["summary"]["best_val_loss"], rtol=0, atol=1e-8):
            raise ValueError("Cached weights do not match validation evidence")
        return result
    print(f"Khôi phục {cfg['exp_id']}, seed {cfg['seed']}; chỉ dùng train/val", flush=True)
    replay = run_experiment({**cfg, "verbose": True},
                            {key: data[key] for key in ("X_tr", "y_tr", "X_val", "y_val")})
    differences = {key: {"original": saved["summary"][key], "replay": replay["summary"][key]}
                   for key in ("best_epoch", "best_val_loss", "val_acc", "val_macro_f1")
                   if not np.isclose(replay["summary"][key], saved["summary"][key], rtol=0, atol=1e-8)}
    if differences:
        # Giữ log gốc và ghi rõ một lần chạy mới; không chọn lại cấu hình hoặc seed.
        if replay["summary"]["best_epoch"] != saved["summary"]["best_epoch"]:
            raise RuntimeError("Recovered best epoch differs from frozen choice; stop before eval")
        from plots import plot_run
        replay["cfg"].update(exp_id=cfg["exp_id"] + "-eval-replay",
                             group="other" if cfg["group"] == "baseline" else "final",
                             description="Part 4: chạy lại đúng cấu hình và seed đã chốt trên val",
                             notes=f"Chạy lại {cfg['exp_id']}; log cũ không lưu trọng số; "
                                   "val khác log cũ; giữ nguyên cấu hình, seed và epoch; không chọn theo eval")
        output = Path(checkpoint_dir).parent
        save_result(replay, str(output / "results"))
        plot_run(replay, str(output / "figures" / f"{replay['cfg']['exp_id']}.png"))
        print("Log chạy lại:", replay["cfg"]["exp_id"], differences, flush=True)
        result = replay
    else:
        result = {**saved, "best_state": replay["best_state"]}
    path.parent.mkdir(parents=True, exist_ok=True)
    torch.save({"cfg": cfg, "result": json.loads(json.dumps(
                   {key: result[key] for key in ("cfg", "history", "summary")})),
                "best_state": replay["best_state"]}, path)
    return result


def score_predictions(repo_root, pred_path, json_path):
    subprocess.run([sys.executable, "scripts/evaluate.py", "--pred", str(Path(pred_path).resolve()),
                    "--out", str(Path(json_path).resolve())], cwd=Path(repo_root).resolve(), check=True)
    scores = json.loads(Path(json_path).read_text(encoding="utf-8"))
    if scores["n_eval"] != 116_203 or np.asarray(scores["confusion_matrix"]).shape != (7, 7):
        raise ValueError("Incomplete official evaluation output")
    return scores


def evaluate_selected(baseline, selected, data, repo_root, output_dir):
    output = Path(output_dir)
    final_eval(baseline["cfg"], baseline, data, str(output / "predictions_eval_baseline.csv"))
    baseline_scores = score_predictions(repo_root, output / "predictions_eval_baseline.csv",
                                        output / "eval_result_baseline.json")
    if selected["cfg"]["exp_id"] == baseline["cfg"]["exp_id"]:
        import shutil
        shutil.copyfile(output / "predictions_eval_baseline.csv", output / "predictions_eval.csv")
        shutil.copyfile(output / "eval_result_baseline.json", output / "eval_result.json")
        final_scores = baseline_scores
    else:
        final_eval(selected["cfg"], selected, data, str(output / "predictions_eval.csv"))
        final_scores = score_predictions(repo_root, output / "predictions_eval.csv", output / "eval_result.json")
    return baseline_scores, final_scores


def error_analysis(scores, figure_path):
    import matplotlib.pyplot as plt

    cm = np.asarray(scores["confusion_matrix"])
    hardest = min(scores["per_class"], key=lambda row: row["f1"])
    cls = hardest["cls"]
    mistakes = cm[cls].copy()
    mistakes[cls] = 0
    confused = int(mistakes.argmax())
    normalized = cm / cm.sum(axis=1, keepdims=True)
    fig, ax = plt.subplots(figsize=(6.5, 4.5))
    im = ax.imshow(normalized, vmin=0, vmax=1, cmap="Blues")
    for i in range(7):
        for j in range(7):
            ax.text(j, i, str(cm[i, j]), ha="center", va="center", fontsize=8,
                    color="white" if normalized[i, j] > .5 else "black")
    ax.set(xticks=range(7), yticks=range(7), xlabel="Predicted label", ylabel="True label",
           title="Final model: eval confusion matrix")
    fig.colorbar(im, ax=ax, label="Fraction within true class")
    fig.tight_layout()
    fig.savefig(figure_path, dpi=150)
    plt.close(fig)
    count = int(mistakes[confused])
    return (f"Lớp khó nhất: **{cls}**, F1={hardest['f1']:.6f}; thường nhầm sang lớp "
            f"**{confused}** ({count}/{hardest['support']} mẫu, "
            f"{count / hardest['support']:.2%}). Hàng là nhãn thật, cột là dự đoán. "
            "Ma trận cho biết số nhầm; đặc trưng giống nhau hoặc ít mẫu học chỉ là giả thuyết, "
            "chưa được kiểm tra bằng thí nghiệm riêng.")


def export_table(output_dir, template_path, baseline, selected, baseline_scores, final_scores):
    output = Path(output_dir)
    results = load_results(str(output / "results"))
    ids = (baseline["cfg"]["exp_id"], selected["cfg"]["exp_id"])
    rows = []
    for result in results:
        exp_id = result["cfg"]["exp_id"]
        scores = baseline_scores if exp_id == ids[0] else final_scores if exp_id == ids[1] else None
        note = ""
        if exp_id == ids[0]:
            note = "Baseline đánh giá eval, seed 1"
        if exp_id == ids[1]:
            note += "; Cấu hình nộp, chọn bằng val; seed=" + str(selected["cfg"]["seed"])
        if not (output / "figures" / f"{exp_id}.png").is_file():
            raise FileNotFoundError(f"Missing experiment figure: {exp_id}")
        rows.append(to_row(result, scores, notes=note.strip("; ")))
    values = [r["summary"]["val_macro_f1"] for r in results if r["cfg"]["group"] == "baseline"]
    noise = 2 * float(np.std(values, ddof=1))
    notes = {
        "baseline": f"Seed 1/2/3; F1 2σ={noise:.6f}. Chỉ seed 1 có eval.",
        "optimizer": "Adam lr=0.003: val F1=0.878221; hơn baseline seed 1 0.021404, dưới 2σ.",
        "hparam": "Batch 128/512/2048: F1=0.852537/0.856817/0.814233; số bước khác nhau.",
        "init": "Zeros chỉ đoán lớp đa số; normal/Xavier gần He, chênh lệch dưới 2σ.",
        "amp": "BF16 CPU: 41.25s/epoch; FP32 baseline 2.04s (log cũ thiếu số luồng). Chưa đo GPU.",
        "final": f"Nộp {ids[1]}, seed {selected['cfg']['seed']}; seed 2/3 chỉ kiểm tra val.",
        "other": "Ba lr SGD+momentum; lần chạy lại baseline giữ riêng, chỉ dòng đó có eval.",
        "loss": "Chưa có JSON thí nghiệm để đối chiếu.",
        "dropout": "Chưa có JSON thí nghiệm để đối chiếu.",
        "clipping": "Chưa có JSON thí nghiệm để đối chiếu.",
    }
    write_xlsx(rows, str(template_path), str(output / "experiments.xlsx"), summary_notes=notes)
    return results


def write_report(output_dir, baseline, selected, baseline_scores, final_scores, error_note):
    """Báo cáo ngắn, lấy số từ JSON kết quả và JSON chấm chính thức."""
    output = Path(output_dir)
    results = {r["cfg"]["exp_id"]: r for r in load_results(str(output / "results"))}
    original_base = results["base-s1"]
    base = baseline["summary"]
    final = selected["summary"]
    seeds = [results[f"base-s{s}"]["summary"] for s in (1, 2, 3)]
    means = {key: float(np.mean([r[key] for r in seeds])) for key in ("val_acc", "val_macro_f1")}
    stds = {key: float(np.std([r[key] for r in seeds], ddof=1)) for key in means}
    noise = 2 * stds["val_macro_f1"]
    small_delta = results['hp-batch128-s1']['summary']['val_macro_f1'] - original_base['summary']['val_macro_f1']
    large_delta = results['hp-batch2048-s1']['summary']['val_macro_f1'] - original_base['summary']['val_macro_f1']
    chosen_original = results['opt-adam-lr3e-3-s1']['summary']
    cfg = selected["cfg"]
    delta = final_scores["macro_f1"] - baseline_scores["macro_f1"]
    rows = []
    for name in ("base-lr0.01-s1", "base-lr0.03-s1", "base-lr0.1-s1",
                 "opt-adam-lr1e-3-s1", "opt-adam-lr3e-3-s1"):
        r = results[name]
        rows.append(f"| `{name}` | {r['cfg']['lr']} | {r['summary']['best_epoch']} | "
                    f"{r['summary']['val_macro_f1']:.6f} |")
    batch_rows = []
    for name in ("hp-batch128-s1", "base-s1", "hp-batch2048-s1"):
        r = results[name]
        steps = int(np.ceil(371847 / r['cfg']['batch'])) * r['cfg']['epochs']
        batch_rows.append(f"| `{name}` | {r['cfg']['batch']} | {steps} | "
                          f"{r['summary']['time_per_epoch_s']:.2f} | {r['summary']['val_macro_f1']:.6f} |")
    class_rows = [f"| {r['cls']} | {r['support']} | {r['precision']:.6f} | "
                  f"{r['recall']:.6f} | {r['f1']:.6f} |" for r in final_scores["per_class"]]
    init_rows = []
    for name in ("init-zeros-s1", "init-normal-s1", "init-xavier-s1"):
        s = results[name]["summary"]
        activation = "/".join(f"{v:.4f}" for v in s["step0_activation_std"])
        init_rows.append(f"| `{name}` | {s['step0_loss']:.6f} | {activation} | {s['val_macro_f1']:.6f} |")
    amp = results["amp-bf16-s1"]["summary"]
    replay_note = ""
    if baseline["cfg"]["exp_id"] != "base-s1" or "eval-replay" in cfg["exp_id"]:
        replay_note = ("Log cũ không lưu trọng số; Part 4 chạy lại đúng cấu hình, seed và epoch. "
                       "Số val của lần chạy lại khác log cũ, nên lưu dòng và ảnh riêng, giữ nguyên log gốc. "
                       "Chưa xác định nguyên nhân chênh lệch; không điều chỉnh cấu hình để khớp số cũ. "
                       "Các nhận xét Part 3 vẫn dùng log gốc; bảng eval bên dưới dùng đúng lần tạo CSV.")
    total = sum(sum(r["history"]["epoch_time_s"]) for r in results.values())
    report = f"""# Báo cáo Lab Day 1 — MSSV 2A202602611

## 1. Thiết lập

Windows, CPU, PyTorch `{torch.__version__}`, Part 3–4 dùng một luồng CPU. Forest CoverType: 464 809 mẫu train gốc, 116 203 eval; tách train/val thành 371 847/92 962 mẫu, phân tầng theo lớp, seed 42. Chuẩn hoá 10 cột số bằng thống kê train; giữ 44 cột nhị phân.

M-base: 54→256→128→7, 47 879 tham số. Baseline `base-s1`: CE (loss phân loại), SGD+momentum 0.9, lr=0.1, batch 512, 20 epoch, He, FP32; dropout và weight decay bằng 0, không clip. Mốc đoán lớp đa số: accuracy val 0.487597 ([Part 2](part2_summary.json)). Chủ đề có JSON đối chiếu: optimizer, batch, init và BF16. Không suy kết luận từ ảnh thiếu JSON.

## 2. Kiểm tra ban đầu và độ nhiễu

| Kiểm tra | Kết quả và nguồn |
|---|---|
| Tham số / đầu ra | 47 879 / (B, 7), notebook Part 1 |
| Loss bước 0 | 2.377572 ở Part 1 seed 42; 2.269062 ở `base-s1` seed 1; ln(7)=1.945910 |
| Học thuộc 20 mẫu | Loss 0.00000725; accuracy 100%, notebook Part 1 |
| Gradient (hướng thay đổi tham số) | Cả 6 tensor có độ lớn khác 0, notebook Part 1 |
| Baseline 3 seed | `base-s1`, `base-s2`, `base-s3` |
| Val accuracy, trung bình ± σ | {means['val_acc']:.6f} ± {stds['val_acc']:.6f} |
| Val macro-F1, trung bình ± σ | {means['val_macro_f1']:.6f} ± {stds['val_macro_f1']:.6f} |

σ là mức dao động giữa các seed; ngưỡng 2σ={noise:.6f} ước lượng từ ba lần chạy. Loss ban đầu cao hơn ln(7) vì điểm số ngẫu nhiên chưa đều, chưa đủ kết luận model lỗi. Công thức bảng so F1 với **trung bình baseline**; phần so seed 1 dưới đây dùng `base-s1`, nên hai chênh lệch khác nhau: cấu hình chọn so trung bình là {chosen_original['val_macro_f1'] - means['val_macro_f1']:+.6f}.

[Ảnh học thuộc 20 mẫu](figures/part1_overfit20.png).

## 3. Kết quả theo chủ đề — chỉ dùng val

### 3.2. Bộ tối ưu và learning rate

Dự đoán trước ở Part 3: Adam điều chỉnh mức cập nhật riêng cho từng trọng số nên có thể giảm loss nhanh hơn; lr=0.003 có thể nhanh hơn 0.001 nhưng dao động hơn.

| exp_id | lr | Epoch val loss thấp nhất | Val macro-F1 |
|---|---:|---:|---:|
{chr(10).join(rows)}

SGD+momentum chọn lr=0.1, Adam chọn 0.003 trong khoảng đã thử. Adam hơn `base-s1` {chosen_original['val_macro_f1'] - original_base['summary']['val_macro_f1']:+.6f}, chưa vượt 2σ={noise:.6f}. SGD lr=0.01 học chậm. Adam 0.001 lại thấp hơn baseline, nên kết luận phụ thuộc lr; chưa chứng minh một optimizer luôn tốt hơn.

![So sánh lr](figures/compare_optimizer_lr.png)

### 3.3. Batch size

Dự đoán trước: batch 128 có nhiều bước hơn nên có thể tăng F1, batch 2048 ít bước hơn nên có thể học chậm.

| exp_id | Batch | Tổng bước | Giây/epoch | Val macro-F1 |
|---|---:|---:|---:|---:|
{chr(10).join(batch_rows)}

Batch 128 không tăng F1 như dự đoán: lệch `base-s1` {small_delta:+.6f}, dưới 2σ. Batch 2048 lệch {large_delta:+.6f}, vượt ngưỡng tham khảo. Số bước ít hơn có thể giải thích kết quả, nhưng chưa thử giữ cùng số bước. Thời gian baseline thiếu số luồng nên chỉ tham khảo. Giữ mọi lô cuối.

![So sánh batch](figures/compare_hparam.png)

### 3.6. BF16 trên CPU — log bổ sung đã có

Log gốc thiếu dự đoán trước; kỳ vọng theo cơ chế là tính ít bit nhanh hơn khi phần cứng hỗ trợ. `amp-bf16-s1`: {amp['time_per_epoch_s']:.2f}s/epoch, accuracy={amp['val_acc']:.6f}, F1={amp['val_macro_f1']:.6f}; `base-s1` FP32: {original_base['summary']['time_per_epoch_s']:.2f}s/epoch, accuracy={original_base['summary']['val_acc']:.6f}, F1={original_base['summary']['val_macro_f1']:.6f}. BF16 chậm hơn, F1 lệch {amp['val_macro_f1'] - original_base['summary']['val_macro_f1']:+.6f}, dưới 2σ. Chi phí đổi kiểu/hỗ trợ CPU là giả thuyết chưa đo riêng. RAM cực đại BF16: {amp['cpu_peak_rss_MB']:.1f} MB; baseline thiếu số RAM tương ứng. Không có GPU; chưa thử FP16.

[Ảnh so sánh precision](figures/compare_amp.png).

### 3.7. Khởi tạo — log bổ sung đã có

Không có dự đoán trước trong log gốc; kỳ vọng theo cơ chế là zeros khó học, He phù hợp ReLU (giữ giá trị dương, đưa giá trị âm về 0). Độ phân tán dưới đây đo sau ba lớp Linear, gồm đầu ra, trên 512 mẫu val đầu.

| exp_id | Loss bước 0 | Độ phân tán 3 lớp | Val macro-F1 |
|---|---:|---|---:|
{chr(10).join(init_rows)}

He seed 1: độ phân tán 0.6662/0.6464/0.5933 (đo lại ở Part 4), loss/F1 `base-s1`=2.269062/{original_base['summary']['val_macro_f1']:.6f}. Zeros có gradient lớp ẩn bằng 0; chỉ bias đầu ra học, accuracy={results['init-zeros-s1']['summary']['val_acc']:.6f}, bằng mốc đa số. Normal làm tín hiệu ban đầu nhỏ nhưng vẫn học được; loss gần ln(7) chưa chứng minh khởi tạo tốt. Normal/Xavier lệch He dưới 2σ; zeros giảm vượt 2σ, mới thử một seed.

[Ảnh so sánh khởi tạo](figures/compare_init.png).

## 4. Đánh giá cuối trên eval

{replay_note}

Cấu hình nộp: **Adam, lr=0.003, batch=512, 20 epoch, seed {cfg['seed']}**, chọn trọng số epoch **{final['best_epoch']}** có val loss thấp nhất; các thành phần khác như baseline. Quyết định gốc nằm trong [part3_selection.json](part3_selection.json), trước khi dùng eval. Cấu hình cuối khác baseline.

| Cấu hình / exp_id | Seed | Val macro-F1 | Eval macro-F1 | Eval accuracy |
|---|---:|---:|---:|---:|
| Baseline `{baseline['cfg']['exp_id']}` | {baseline['cfg']['seed']} | {base['val_macro_f1']:.6f} | {baseline_scores['macro_f1']:.6f} | {baseline_scores['accuracy']:.6f} |
| Cuối `{cfg['exp_id']}` | {cfg['seed']} | {final['val_macro_f1']:.6f} | {final_scores['macro_f1']:.6f} | {final_scores['accuracy']:.6f} |

Nguồn eval: [baseline](eval_result_baseline.json), [cấu hình cuối](eval_result.json), đều do `scripts/evaluate.py` tạo, n_eval=116203. Eval F1 đổi {delta:+.6f}; {'chưa vượt' if abs(delta) <= noise else 'vượt'} ngưỡng val 2σ={noise:.6f}, nhưng **chưa đo nhiễu eval** vì chỉ chấm seed 1. Không thay tham số theo kết quả này. F1 eval–val của cấu hình cuối lệch {final_scores['macro_f1'] - final['val_macro_f1']:+.6f}.

### 4.1. Lỗi theo lớp

Precision là tỷ lệ đúng trong các mẫu được đoán thuộc lớp; recall là tỷ lệ tìm đúng trong mẫu thật của lớp. Support là số mẫu thật; F1 kết hợp precision và recall.

| Lớp | Support | Precision | Recall | F1 |
|---|---:|---:|---:|---:|
{chr(10).join(class_rows)}

{error_note} Lớp 3 chỉ có {final_scores['per_class'][3]['support']} mẫu nhưng F1={final_scores['per_class'][3]['f1']:.6f}, cao hơn lớp 4; ít mẫu không giải thích được toàn bộ khó khăn. Có thể thử tăng trọng số loss của lớp khó trên train và chọn bằng val; đây là đề xuất chưa thử, không sửa mô hình sau eval.

![Ma trận nhầm lẫn, số trong ô là số mẫu](figures/compare_eval_confusion.png)

## 5. Câu hỏi dẫn dắt

- **Câu 1 — optimizer:** Adam 0.003 cao nhất trong lr đã thử; nếu chỉ dùng Adam 0.001 thì thứ hạng đảo (bảng 3.2). Phải chỉnh lr cho mỗi bộ và xét nhiễu.
- **Câu 4 — precision:** BF16 CPU chậm hơn (`amp-bf16-s1`); chưa đo riêng chi phí đổi kiểu hoặc GPU.
- **Câu 5 — khởi tạo:** zeros chặn gradient lớp ẩn. He bù tín hiệu ReLU bỏ đi; Xavier cân bằng tín hiệu vào/ra. Khác biệt quan trọng khi nhiều lớp làm tín hiệu suy giảm; log này chưa chứng minh He vượt Xavier ổn định.
- **Câu 6 — ba kiểm tra đầu:** (1) kiểm tra nhãn 0..6, ghép mẫu–nhãn, chuẩn hoá và đầu ra (B,7), loại lỗi dữ liệu/loss; (2) xem gradient khác 0, hữu hạn và trọng số có đổi sau cập nhật, kiểm tra model thực sự học (`init-zeros-s1` là ví dụ lỗi); (3) thử học thuộc 20 mẫu, tắt dropout và khảo sát vài lr trên train/val, phân biệt lỗi pipeline với tốc độ học sai. Part 1 đạt loss 0.00000725.

## 6. Hạn chế và điều bất ngờ

Batch nhỏ không cải thiện như dự đoán; BF16 CPU chậm. Các chủ đề bổ sung thiếu dự đoán trước được lưu; số seed ít, các batch khác tổng bước và số đo thời gian cũ thiếu thông tin môi trường. Seed 2/3 của cấu hình Adam (`final-check-s2`, `final-check-s3`) chỉ có val, không dùng để chọn seed nộp. Giả thuyết về nguyên nhân nhầm lớp và sai lệch khi chạy lại chưa được kiểm tra. Nếu có thêm thời gian: so batch với cùng số bước, nhiều seed và khảo sát đặc trưng của lớp khó; thực hiện trên train/val.

## 7. Phụ lục

File chính: `code/lab.ipynb`, `code/train.py`, `code/results_table.py`, `predictions_eval.csv`, `eval_result.json`, `experiments.xlsx`, `REPORT.md`, `results/`, `figures/`. Baseline giữ riêng `predictions_eval_baseline.csv`, `eval_result_baseline.json`; [manifest](part4_manifest.json) ghi đúng seed, epoch và exp_id tạo file nộp. JSON lịch sử chỉ chứa cấu hình, lịch sử, tóm tắt; trọng số nằm riêng trong `checkpoints/` (bị gitignore).

Tổng thời gian các epoch trong JSON: khoảng {total / 60:.1f} phút, không gồm đọc dữ liệu, kiểm tra Part 1 và chấm. Bảng giữ bốn sheet/cột/công thức của mẫu; Excel tính lại trước khi kiểm tra.
"""
    path = output / "REPORT.md"
    path.write_text(report, encoding="utf-8")
    return path


def verify_submission(output_dir, baseline_id, final_id, require_calculated=False):
    import re
    import pandas as pd
    import openpyxl

    output = Path(output_dir)
    repo = output.resolve().parent
    with np.load(repo / "data" / "processed" / "eval.npz") as data:
        expected = data["row_id"]
    for suffix in ("", "_baseline"):
        pred = pd.read_csv(output / f"predictions_eval{suffix}.csv")
        assert list(pred.columns) == ["row_id", "pred"]
        assert len(pred) == 116_203 and pred.row_id.is_unique
        assert np.array_equal(pred.row_id.to_numpy(), expected), "Prediction order no longer matches eval.npz"
        assert pred.pred.dtype.kind in "iu" and pred.pred.between(0, 6).all()
        score = json.loads((output / f"eval_result{suffix}.json").read_text(encoding="utf-8"))
        assert score["n_eval"] == 116_203
        assert np.asarray(score["confusion_matrix"]).sum() == 116_203
        assert len(score["per_class"]) == 7
    workbook = openpyxl.load_workbook(output / "experiments.xlsx")
    assert workbook.sheetnames == ["Legend", "Experiments", "Seeds", "Summary"]
    template = openpyxl.load_workbook(repo / "templates" / "experiment_table_template.xlsx")
    for name in ("Experiments", "Seeds", "Summary"):
        assert [c.value for c in workbook[name][1]] == [c.value for c in template[name][1]]
    for sheet in template:
        for cells in sheet:
            for cell in cells:
                if cell.data_type == "f":
                    assert workbook[sheet.title][cell.coordinate].data_type == "f", cell.coordinate
    experiments = workbook["Experiments"]
    headers = [c.value for c in experiments[1]]
    rows = {r[0].value: dict(zip(headers, [c.value for c in r]))
            for r in list(experiments)[1:] if r[0].value}
    results = load_results(str(output / "results"))
    assert len(rows) == len(results)
    assert set(rows) == {r["cfg"]["exp_id"] for r in results}
    scored_ids = {key for key, row in rows.items() if row["eval_acc"] is not None or row["eval_macro_f1"] is not None}
    assert scored_ids == {baseline_id, final_id}
    for exp_id, suffix in ((baseline_id, "_baseline"), (final_id, "")):
        score = json.loads((output / f"eval_result{suffix}.json").read_text(encoding="utf-8"))
        assert abs(rows[exp_id]["eval_acc"] - score["accuracy"]) < 1e-12
        assert abs(rows[exp_id]["eval_macro_f1"] - score["macro_f1"]) < 1e-12
    for exp_id, row in rows.items():
        assert row["figure_file"] == f"figures/{exp_id}.png" and (output / row["figure_file"]).exists()
    report = (output / "REPORT.md").read_text(encoding="utf-8")
    for target in re.findall(r"\]\(([^)]+)\)", report):
        assert (output / target).exists(), target
    if require_calculated:
        cached = openpyxl.load_workbook(output / "experiments.xlsx", data_only=True)
        for sheet in cached:
            for cells in sheet:
                for cell in cells:
                    assert cell.data_type != "e", f"Excel error: {sheet.title}!{cell.coordinate}={cell.value}"
        assert abs(cached["Seeds"]["C10"].value - 2 * np.std(
            [rows[f"base-s{s}"]["val_macro_f1"] for s in (1, 2, 3)], ddof=1)) < 1e-12
        assert cached["Summary"]["D13"].value == 4
    print(f"Đối chiếu OK: {len(rows)} thí nghiệm; 2 dòng có eval; CSV, JSON, ảnh, bảng và đường dẫn báo cáo hợp lệ.")
